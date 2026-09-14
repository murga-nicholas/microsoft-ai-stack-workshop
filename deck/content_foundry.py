"""Microsoft Foundry architecture, jobs and recorded modes."""

from __future__ import annotations

from typing import TYPE_CHECKING

import components as C
import diagrams as D
import research_facts as R
import theme as T
from theme import P, plain, text_box, write

if TYPE_CHECKING:
    from facts import Facts
    from pptx.presentation import Presentation
    from pptx.slide import Slide

U = 914400
HALF = (T.CONTENT_W - 274320) // 2
RIGHT = T.MARGIN_L + HALF + 274320
CASE_Y = int(7.55 * U)


def _text(
    slide: Slide,
    text: str,
    *,
    x: int = T.MARGIN_L,
    y: int,
    w: int = T.CONTENT_W,
    h: int = 548640,
    size: float = 14,
    colour: str = P.muted,
    bold: bool = False,
) -> None:
    """Place a readable body sentence on the shared grid."""
    box = text_box(slide, x=x, y=y, w=w, h=h)
    write(box, plain(text), size=size, colour=colour, bold=bold)


def _notes(slide: Slide, *urls: str) -> None:
    """Keep research provenance in the speaker notes."""
    slide.notes_slide.notes_text_frame.text = "Sources:\n" + "\n".join(urls)


def slide_section(prs: Presentation, facts: Facts) -> None:
    """Introduce the services covered in the Foundry segment."""
    del facts
    T.section_slide(
        prs,
        index="05",
        eyebrow="microsoft foundry",
        title_lines=["The cloud takes", "real responsibility"],
        deck="Models, agents and knowledge share a project. Each service "
        "has a distinct runtime and governance boundary.",
        bullets=(
            "The catalog, project endpoint and openai package",
            "Prompt agents, hosted agents and Foundry Local",
            "Foundry Local or Ollama?",
            "Azure AI Search and Foundry IQ",
            "Guardrails, identity and evaluation",
        ),
        number=25,
        accent=P.sky,
    )


def slide_models(prs: Presentation, facts: Facts) -> None:
    """Connect the model catalog to the project and model clients."""
    del facts
    slide = T.content_slide(
        prs,
        eyebrow="models and project clients",
        title="One catalog, one project endpoint, the openai package",
        deck="The catalog helps you choose. AIProjectClient manages the "
        "project, and the openai package calls its model endpoint.",
        number=26,
        accent=P.sky,
    )
    C.run_strip(slide, "azure-ai-projects", "foundry-models")
    C.screenshot(
        slide,
        str(T.ASSETS / "foundry_model_catalog.jpg"),
        x=T.MARGIN_L,
        y=T.Y_BODY,
        w=int(4.6 * U),
        h=int(1.62 * U),
        caption=f"Foundry model catalog - {R.MODEL_CATALOG_COUNT} models "
        f"visible in this capture, {R.CAPTURE_DATE}",
    )
    D.draw(
        slide,
        D.Schema(
            nodes=[
                D.Node(
                    "project",
                    "AIProjectClient",
                    0,
                    0,
                    13,
                    2,
                    colour=P.sky,
                    mono=True,
                ),
                D.Node(
                    "ops",
                    "deployments / connections",
                    14.2,
                    0,
                    18.8,
                    2,
                    kind="note",
                ),
                D.Node(
                    "get",
                    "get_openai_client()",
                    0,
                    3,
                    13,
                    2,
                    colour=P.blue,
                    mono=True,
                ),
                D.Node(
                    "resp", "Responses API", 14.2, 3, 18.8, 2, colour=P.blue
                ),
                D.Node(
                    "af",
                    "FoundryChatClient",
                    0,
                    6,
                    13,
                    2,
                    colour=P.teal,
                    mono=True,
                ),
                D.Node(
                    "v1",
                    "<project endpoint>/openai/v1",
                    14.2,
                    6,
                    18.8,
                    2,
                    kind="note",
                    mono=True,
                ),
            ],
            edges=[
                D.Edge("project", "ops", side="h"),
                D.Edge("project", "get", side="v"),
                D.Edge("get", "resp", side="h"),
                D.Edge("af", "get", side="v"),
                D.Edge("resp", "v1", side="v"),
            ],
        ),
        x=RIGHT,
        y=T.Y_BODY,
    )
    _text(
        slide,
        "Entra audience: https://ai.azure.com/.default",
        x=RIGHT,
        y=int(5.16 * U),
        w=HALF,
        h=365760,
    )
    _text(
        slide,
        f"SDK migration: azure-ai-inference (retired "
        f"{R.INFERENCE_RETIRED}) → openai\nAzureOpenAI() → "
        "OpenAI(base_url) · azure-ai-projects 1.x → 2.x",
        x=T.MARGIN_L + int(4.87 * U),
        y=T.Y_BODY + int(0.13 * U),
        w=int(3.2 * U),
        h=int(1.82 * U),
    )
    _text(
        slide,
        "Developer maintains a meeting summarizer (illustrative)",
        y=int(5.68 * U),
        w=int(11.8 * U),
        h=int(0.3 * U),
        size=16,
        colour=P.ink_soft,
        bold=True,
    )
    _text(
        slide,
        "Legacy",
        y=int(6.2 * U),
        w=int(0.85 * U),
        h=int(0.3 * U),
        colour=P.coral,
        bold=True,
    )
    D.draw(
        slide,
        D.Schema(
            nodes=[
                D.Node(
                    "transcript", "transcript", 0, 0, 4.8, 2.8, kind="note"
                ),
                D.Node(
                    "complete",
                    "ChatCompletionsClient.complete()",
                    6.4,
                    0,
                    18,
                    2.8,
                    colour=P.coral,
                    mono=True,
                    sub="azure-ai-inference",
                ),
                D.Node("draft", "summary draft", 26, 0, 5.8, 2.8, kind="note"),
            ],
            edges=[
                D.Edge("transcript", "complete", side="h"),
                D.Edge("complete", "draft", side="h"),
            ],
        ),
        x=T.MARGIN_L + int(1.0 * U),
        y=int(6.06 * U),
    )
    retired = T.rounded(
        slide,
        x=T.MARGIN_L + int(9.25 * U),
        y=int(6.16 * U),
        w=int(2.65 * U),
        h=int(0.36 * U),
        fill=P.wash,
        line=P.coral,
        radius=50000,
    )
    retired.text_frame.margin_left = 91440
    retired.text_frame.margin_right = 91440
    retired.text_frame.margin_top = 45720
    retired.text_frame.margin_bottom = 0
    write(
        retired,
        plain(f"SDK retired {R.INFERENCE_RETIRED}"),
        size=14,
        colour=P.coral,
        bold=True,
    )
    command = text_box(
        slide,
        x=T.MARGIN_L + int(12.18 * U),
        y=int(5.7 * U),
        w=T.CONTENT_W - int(12.18 * U),
        h=int(0.3 * U),
    )
    write(
        command,
        plain("uv run msai-demo azure-ai-inference"),
        size=14,
        colour=P.ink_soft,
        bold=True,
        font=T.FONT_CODE,
    )
    _text(
        slide,
        "local_contract - validates the legacy request shape and decodes "
        "a synthetic response; no SDK or model runs",
        x=T.MARGIN_L + int(12.18 * U),
        y=int(6.06 * U),
        w=T.CONTENT_W - int(12.18 * U),
        h=int(0.73 * U),
        size=14,
    )
    _text(
        slide,
        "Current",
        y=int(7.0 * U),
        w=int(0.85 * U),
        h=int(0.3 * U),
        colour=P.teal,
        bold=True,
    )
    D.draw(
        slide,
        D.Schema(
            nodes=[
                D.Node(
                    "transcript", "transcript", 0, 0, 4.8, 2.8, kind="note"
                ),
                D.Node(
                    "openai",
                    "OpenAI(base_url=<project endpoint>/openai/v1)",
                    6.4,
                    0,
                    23,
                    2.8,
                    colour=P.teal,
                    mono=True,
                ),
                D.Node(
                    "responses",
                    "responses.create()",
                    31,
                    0,
                    9.6,
                    2.8,
                    colour=P.blue,
                    mono=True,
                ),
                D.Node(
                    "draft", "summary draft", 42.2, 0, 5.8, 2.8, kind="note"
                ),
            ],
            edges=[
                D.Edge("transcript", "openai", side="h"),
                D.Edge("openai", "responses", side="h"),
                D.Edge("responses", "draft", side="h"),
            ],
        ),
        x=T.MARGIN_L + int(1.0 * U),
        y=int(6.86 * U),
    )
    C.use_case(
        slide,
        who="Analyst picks a model for meeting summaries",
        steps=(
            "shortlist in catalog",
            "same transcript",
            "compare omissions / latency",
        ),
        result="a justified choice.",
        x=T.MARGIN_L,
        y=CASE_Y,
        w=HALF,
        accent=P.sky,
    )
    C.use_case(
        slide,
        who="Developer needs a deployment name",
        steps=("authenticate to project", "list deployments", "choose one"),
        result="configuration for the app.",
        x=RIGHT,
        y=CASE_Y,
        w=HALF,
        accent=P.teal,
    )
    _notes(
        slide,
        "https://learn.microsoft.com/en-us/azure/foundry/how-to/develop/sdk-overview",
        "https://learn.microsoft.com/en-us/azure/foundry/how-to/navigate-from-classic",
    )


def slide_agent_service(prs: Presentation, facts: Facts) -> None:
    """Compare declarative agents with managed hosting of code."""
    del facts
    slide = T.content_slide(
        prs,
        eyebrow="managed agents",
        title="Prompt agents or hosted agents: who runs the code?",
        deck=R.AGENT_SERVICE_DEFINITION,
        number=27,
        accent=P.sky,
    )
    C.run_strip(slide, "foundry-agent-service")
    for x, prefix, definition, accent in (
        (T.MARGIN_L, "Prompt agent", "instructions + model + tools", P.sky),
        (RIGHT, "Hosted agent", "your framework + agent code", P.teal),
    ):
        _text(
            slide,
            prefix,
            x=x,
            y=T.Y_BODY,
            w=HALF,
            size=18,
            bold=True,
            colour=accent,
        )
        D.draw(
            slide,
            D.Schema(
                nodes=[
                    D.Node("define", definition, 0, 0, 14, 3.2, colour=accent),
                    D.Node(
                        "runtime",
                        "Foundry runtime",
                        18,
                        0,
                        14,
                        3.2,
                        colour=P.blue,
                    ),
                    D.Node(
                        "app",
                        "your application",
                        0,
                        6,
                        14,
                        3.2,
                        kind="actor",
                        colour=accent,
                    ),
                    D.Node(
                        "endpoint",
                        "managed endpoint",
                        18,
                        6,
                        14,
                        3.2,
                        colour=P.blue,
                    ),
                ],
                edges=[
                    D.Edge("define", "runtime", side="h"),
                    D.Edge("runtime", "endpoint", side="v"),
                    D.Edge("app", "endpoint", side="h"),
                ],
            ),
            x=x,
            y=T.Y_BODY + 502920,
        )
    _text(
        slide,
        "Declarative definition. Foundry hosts the model/tool loop.",
        x=T.MARGIN_L,
        y=int(6.02 * U),
        w=HALF,
    )
    _text(
        slide,
        "Agent Framework, LangGraph, Semantic Kernel or your code. "
        "Foundry provides scaling, identity and observability.",
        x=RIGHT,
        y=int(6.02 * U),
        w=HALF,
        h=640080,
    )
    _text(
        slide,
        "Offline mode=local_contract exercises requests and replies. "
        "Live hosting needs a project and role.",
        y=int(6.9 * U),
        h=411480,
        colour=P.coral,
    )
    C.use_case(
        slide,
        who="HR administrator needs policy Q&A",
        steps=("pick model", "write instructions", "attach knowledge", "test"),
        result="an assistant with no runtime to operate.",
        x=T.MARGIN_L,
        y=CASE_Y,
        w=HALF,
        accent=P.sky,
    )
    C.use_case(
        slide,
        who="Python team hosts its log analyzer",
        steps=("package agent code", "deploy container", "call endpoint"),
        result="their logic, managed hosting.",
        x=RIGHT,
        y=CASE_Y,
        w=HALF,
        accent=P.teal,
    )
    _notes(slide, R.FOUNDRY_AGENT_SERVICE_URL)


def slide_foundry_local(prs: Presentation, facts: Facts) -> None:
    """Show the native runtime and distinguish the demo adapter."""
    del facts
    slide = T.content_slide(
        prs,
        eyebrow="local inference",
        title="Foundry Local runs the model inside your app",
        deck="Foundry Local is an embedded native runtime for on-device "
        "inference, with SDKs for C#, JavaScript, Python and Rust.",
        number=28,
        accent=P.sky,
    )
    C.run_strip(slide, "foundry-local")
    D.draw(
        slide,
        D.Schema(
            nodes=[
                D.Node("app", "your app", 0, 0, 8, 3.2, colour=P.teal),
                D.Node("sdk", "language SDK", 10, 0, 9, 3.2, colour=P.sky),
                D.Node(
                    "core",
                    "Foundry Local Core API",
                    21,
                    0,
                    17,
                    3.2,
                    colour=P.blue,
                ),
                D.Node("onnx", "ONNX Runtime", 40, 0, 12, 3.2, colour=P.blue),
                D.Node(
                    "providers",
                    "execution providers",
                    54,
                    0,
                    12,
                    3.2,
                    colour=P.violet,
                ),
                D.Node(
                    "rest",
                    "optional OpenAI REST endpoint",
                    0,
                    6,
                    18,
                    3,
                    kind="note",
                ),
                D.Node(
                    "cache",
                    "local model cache",
                    21,
                    6,
                    17,
                    3,
                    kind="state",
                    colour=P.sky,
                ),
                D.Node(
                    "catalog",
                    "Foundry Catalog",
                    47,
                    6,
                    17,
                    3,
                    kind="actor",
                    colour=P.sky,
                ),
            ],
            edges=[
                D.Edge("app", "sdk", side="h"),
                D.Edge("sdk", "core", side="h"),
                D.Edge("core", "onnx", side="h"),
                D.Edge("onnx", "providers", side="h"),
                D.Edge("rest", "core", side="h"),
                D.Edge("core", "cache", side="v"),
                D.Edge(
                    "catalog",
                    "cache",
                    "first-use download",
                    side="h",
                    dashed=True,
                ),
            ],
        ),
        x=T.MARGIN_L,
        y=T.Y_BODY,
    )
    _text(
        slide,
        R.LOCAL_HARDWARE_NOTE,
        y=int(5.62 * U),
        h=640080,
        size=15,
    )
    _text(
        slide,
        R.LOCAL_NETWORK_NOTE,
        y=int(6.4 * U),
    )
    _text(
        slide,
        "SDK traps: foundry_local became foundry_local_sdk at 2.x; "
        "agent-framework-foundry-local pins 0.5.1. This repo uses the "
        "OpenAI-compatible endpoint.",
        y=int(7.02 * U),
        h=502920,
    )
    C.use_case(
        slide,
        who="Field technician needs notes summarised offline",
        steps=(
            "download model once",
            "load in-process",
            "summarise local text",
        ),
        result="an on-device draft, no cloud call.",
        x=T.MARGIN_L,
        y=int(7.72 * U),
        w=T.CONTENT_W,
        accent=P.sky,
    )
    C.source_note(
        slide,
        "Offline: labelled synthetic reply. "
        "Live without runtime: blocked / service_unavailable. "
        "Source: Microsoft Learn, Foundry Local architecture.",
    )
    _notes(slide, R.FOUNDRY_LOCAL_ARCHITECTURE_URL)


def slide_local_comparison(prs: Presentation, facts: Facts) -> None:
    """Explain where Foundry Local and Ollama overlap and differ."""
    del facts
    slide = T.content_slide(
        prs,
        eyebrow="local runtime choice",
        title="Foundry Local or Ollama? Overlap, not a drop-in swap",
        deck="Both can run local models. The app integration, model formats "
        "and runtime ownership differ.",
        number=29,
        accent=P.sky,
    )
    C.run_strip(slide, "foundry-local")
    C.table(
        slide,
        ("", "Foundry Local", "Ollama"),
        R.LOCAL_COMPARISON_ROWS,
        x=T.MARGIN_L,
        y=T.Y_BODY,
        w=T.CONTENT_W,
        widths=(1.3, 4, 4),
        row_h=411480,
        size=14,
    )
    _text(
        slide,
        "For local development they overlap. Ship AI inside a "
        "Windows, macOS or Linux app: Foundry Local. Tinker with community "
        "models behind a local server: Ollama. For high-throughput shared "
        "serving, prefer a serving stack such as vLLM.",
        y=int(6.7 * U),
        h=685800,
        size=15,
        colour=P.ink_soft,
        bold=True,
    )
    C.use_case(
        slide,
        who="Field technician uses Foundry Local offline",
        steps=("download once", "embed runtime", "summarise notes"),
        result="an on-device draft.",
        x=T.MARGIN_L,
        y=CASE_Y,
        w=HALF,
        accent=P.sky,
    )
    C.use_case(
        slide,
        who="Python developer experiments with Ollama",
        steps=(
            "ollama pull a model",
            "run server",
            "call /v1/chat/completions",
        ),
        result="a local development endpoint.",
        x=RIGHT,
        y=CASE_Y,
        w=HALF,
        accent=P.teal,
    )
    C.source_note(
        slide,
        "Sources: Microsoft Learn, Foundry Local architecture and "
        "FAQ; Ollama documentation and its OpenAI compatibility page.",
    )
    _notes(
        slide,
        R.FOUNDRY_LOCAL_ARCHITECTURE_URL,
        R.OLLAMA_COMPATIBILITY_URL,
        "https://docs.ollama.com/import",
        "https://docs.ollama.com/faq",
        "https://github.com/microsoft/Foundry-Local",
    )
    slide.notes_slide.notes_text_frame.text += (
        "\nThe verdict is the workshop's architecture recommendation. "
        "Ollama supports concurrent requests and network serving; "
        "the recommendation does not imply technical inability."
    )


def slide_knowledge(prs: Presentation, facts: Facts) -> None:
    """Show Search as the retrieval engine beneath Foundry IQ."""
    del facts
    slide = T.content_slide(
        prs,
        eyebrow="knowledge",
        title="AI Search is the engine; Foundry IQ is the knowledge layer",
        deck=R.FOUNDRY_IQ_DEFINITION,
        number=30,
        accent=P.sky,
    )
    C.run_strip(slide, "ai-search", "foundry-iq")
    D.draw(
        slide,
        D.Schema(
            nodes=[
                D.Node("query", "your query", 0, 0, 12, 3, colour=P.teal),
                D.Node(
                    "search", "Azure AI Search", 18, 0, 18, 3, colour=P.blue
                ),
                D.Node(
                    "passages",
                    "ranked passages + sources",
                    45,
                    0,
                    20,
                    3,
                    kind="note",
                ),
                D.Node("agent", "agent request", 0, 8, 12, 3, colour=P.teal),
                D.Node(
                    "iq",
                    "Foundry IQ knowledge base",
                    18,
                    8,
                    18,
                    3,
                    colour=P.sky,
                ),
                D.Node(
                    "context",
                    "grounded context + citations",
                    45,
                    8,
                    20,
                    3,
                    kind="note",
                ),
            ],
            edges=[
                D.Edge("query", "search", "hybrid query", side="h"),
                D.Edge("search", "passages", side="h"),
                D.Edge("agent", "iq", side="h"),
                D.Edge("iq", "context", "plans + merges", side="h"),
                D.Edge("iq", "search", "built on", side="v", dashed=True),
            ],
        ),
        x=T.MARGIN_L,
        y=T.Y_BODY + 91440,
    )
    C.table(
        slide,
        ("You configure", "Azure AI Search", "Foundry IQ"),
        (
            (
                "Retrieval",
                "Index, schema, refresh and ranking",
                "Knowledge sources and retrieval settings",
            ),
            (
                "Permissions",
                "Security filters or supported ACL / RBAC",
                "Supported source ACL sync + query-time trimming",
            ),
        ),
        x=T.MARGIN_L,
        y=int(6.05 * U),
        w=T.CONTENT_W,
        widths=(1.4, 3.8, 3.8),
        row_h=411480,
        size=14,
    )
    C.use_case(
        slide,
        who="Support engineer needs manual passages",
        steps=(
            "index manuals",
            "hybrid query",
            "ranked passages with sources",
        ),
        result="relevant evidence for the answer.",
        x=T.MARGIN_L,
        y=CASE_Y,
        w=HALF,
        accent=P.blue,
    )
    C.use_case(
        slide,
        who="Employee needs a cross-document policy answer",
        steps=(
            "sources + permissions",
            "agent asks knowledge base",
            "cited context",
        ),
        result="a grounded policy answer.",
        x=RIGHT,
        y=CASE_Y,
        w=HALF,
        accent=P.sky,
    )
    C.source_note(
        slide,
        f"Foundry IQ: {R.FOUNDRY_IQ_STATUS}. "
        "Configure permission trimming before content reaches the model.",
    )
    _notes(
        slide,
        "https://learn.microsoft.com/en-us/azure/foundry/concepts/general-availability",
        "https://learn.microsoft.com/en-us/azure/foundry/agents/concepts/what-is-foundry-iq",
        "https://learn.microsoft.com/en-us/azure/search/search-what-is-azure-search",
    )


def slide_guardrails(prs: Presentation, facts: Facts) -> None:
    """Separate service controls from the recorded local policy."""
    run = facts.demo_run("foundry-guardrails")
    over = facts.guardrail_case("over_budget")
    slide = T.content_slide(
        prs,
        eyebrow="safety",
        title="Classifying a prompt is not the same as intervening",
        deck="Content Safety returns severity scores. Guardrails use "
        "classifiers and configured actions at input, tool and output "
        "boundaries.",
        number=31,
        accent=P.coral,
    )
    C.run_strip(slide, "foundry-guardrails", "content-safety")
    nodes = []
    labels = (
        "user input",
        "model / agent",
        "tool call",
        "tool response",
        "output",
    )
    for index, label in enumerate(labels):
        nodes.append(
            D.Node(
                str(index),
                label,
                index * 10,
                0,
                8.5,
                3,
                colour=P.coral if index in (2, 3) else P.blue,
            )
        )
    D.draw(
        slide,
        D.Schema(
            nodes=nodes,
            edges=[D.Edge(str(i), str(i + 1), side="h") for i in range(4)],
        ),
        x=T.MARGIN_L,
        y=T.Y_BODY,
    )
    _text(
        slide,
        "Input and output: "
        + R.GUARDRAILS_MODEL_AGENT_STATUS
        + ".\nTool call and tool response interception: "
        + R.GUARDRAILS_TOOL_STATUS
        + ".",
        y=T.Y_BODY + 685800,
        w=T.COL_MAIN_W,
        h=685800,
        size=15,
    )
    _text(
        slide,
        "Agent guardrails apply to Foundry Agent Service agents.",
        y=T.Y_BODY + 1417320,
        w=T.COL_MAIN_W,
        h=411480,
        bold=True,
        colour=P.ink_soft,
    )
    C.table(
        slide,
        ("Recorded local point", "Action", "Matched controls"),
        tuple(
            (
                str(d["point"]).replace("_", " "),
                str(d["action"]),
                ", ".join(d["matched_controls"]) or "-",
            )
            for d in over["decisions"]
        ),
        x=T.MARGIN_L,
        y=T.Y_BODY + 1874520,
        w=T.COL_MAIN_W,
        widths=(2, 1, 4),
        row_h=274320,
        size=14,
    )
    C.rail(
        slide,
        title=str(run["mode"]).replace("_", " ").upper(),
        lines=(
            "Synthetic signals; application-owned budget gate.",
            f"Recorded: {run['headline']}",
            "Foundry enforcement needs a separate live validation.",
        ),
        accent=P.coral,
        h=3383280,
    )
    _text(
        slide,
        "Content Safety: classify text into severities; your app "
        "applies the threshold. Content Moderator retires "
        f"{R.CONTENT_MODERATOR_RETIRES}.",
        y=int(7.17 * U),
        h=411480,
    )
    C.use_case(
        slide,
        who="Community moderator flags harmful posts",
        steps=("submit text", "read severity scores", "apply threshold"),
        result="allow, block or review.",
        x=T.MARGIN_L,
        y=int(7.75 * U),
        w=HALF,
        accent=P.blue,
    )
    C.use_case(
        slide,
        who="Assistant owner blocks risky actions",
        steps=(
            "attach guardrail",
            "check input / tool / output",
            "configured action",
        ),
        result="the risky call never runs.",
        x=RIGHT,
        y=int(7.75 * U),
        w=HALF,
        accent=P.coral,
    )
    C.source_note(
        slide,
        "Local example blocked the simulated business "
        "action. Service guardrails use Content Safety classifiers.",
    )
    _notes(
        slide,
        "https://learn.microsoft.com/en-us/azure/foundry/concepts/general-availability",
        "https://learn.microsoft.com/en-us/azure/foundry/guardrails/guardrails-overview",
        "https://learn.microsoft.com/en-us/azure/ai-services/content-moderator/overview",
    )


def slide_identity(prs: Presentation, facts: Facts) -> None:
    """Teach tokens, role assignments and agent-specific governance."""
    run = facts.identity
    slide = T.content_slide(
        prs,
        eyebrow="identity",
        title="A token proves who you are. It does not grant access.",
        deck="Microsoft Entra ID authenticates the app; resource roles "
        "authorize it. Entra Agent ID extends Entra ID with agent identities.",
        number=32,
        accent=P.amber,
    )
    C.run_strip(slide, "azure-identity", "entra-agent-id")
    D.draw(
        slide,
        D.Schema(
            nodes=[
                D.Node("app", "app identity", 0, 0, 12, 3, colour=P.amber),
                D.Node("token", "Entra token", 17, 0, 12, 3, colour=P.teal),
                D.Node("role", "resource role", 34, 0, 12, 3, colour=P.blue),
                D.Node("reply", "allowed / 403", 51, 0, 13, 3, kind="note"),
                D.Node("bp", "agent blueprint", 0, 6, 12, 3, colour=P.amber),
                D.Node("id", "agent identity", 17, 6, 12, 3, colour=P.sky),
                D.Node("grant", "scoped access", 34, 6, 12, 3, colour=P.blue),
                D.Node("review", "review / revoke", 51, 6, 13, 3, kind="note"),
            ],
            edges=[
                D.Edge("app", "token", side="h"),
                D.Edge("token", "role", side="h"),
                D.Edge("role", "reply", side="h"),
                D.Edge("bp", "id", side="h", dashed=True),
                D.Edge("id", "grant", side="h", dashed=True),
                D.Edge("grant", "review", side="h", dashed=True),
            ],
        ),
        x=T.MARGIN_L,
        y=T.Y_BODY,
    )
    tokens = run.get("tokens") or []
    arms = run.get("arm") or {}
    subscriptions = arms.get("subscriptions")
    outcomes = (
        "; ".join(
            str(t["scope"]).removeprefix("https://").removesuffix("/.default")
            + (": acquired" if t.get("acquired") else ": not acquired")
            for t in tokens
        )
        or "no token calls recorded"
    )
    _text(
        slide,
        f"Recorded mode={run['mode']}: {outcomes}.",
        y=int(5.65 * U),
        h=685800,
    )
    _text(
        slide,
        "GET /subscriptions: "
        + (
            f"{len(subscriptions)} subscriptions"
            if isinstance(subscriptions, list)
            else "not called"
        )
        + ". Token values are never printed.",
        y=int(6.35 * U),
        h=365760,
    )
    _text(
        slide,
        "The workshop principal is an ordinary app registration, "
        "not an Agent ID. The dashed path models what a tenant admin "
        "must create first; this repository did not provision it.",
        y=int(6.85 * U),
        colour=P.coral,
    )
    C.use_case(
        slide,
        who="Developer needs Search access for an app",
        steps=("acquire token", "call Search", "role decides"),
        result="allowed, or a clear 403.",
        x=T.MARGIN_L,
        y=CASE_Y,
        w=HALF,
        accent=P.teal,
    )
    C.use_case(
        slide,
        who="Tenant admin needs an accountable agent",
        steps=(
            "blueprint + identity",
            "grant scoped access",
            "review or revoke",
        ),
        result="agent access governed separately from people.",
        x=RIGHT,
        y=CASE_Y,
        w=HALF,
        accent=P.amber,
    )
    _notes(
        slide,
        "https://learn.microsoft.com/en-us/entra/agent-id/identity-platform/what-is-agent-id",
    )


def slide_evaluation(prs: Presentation, facts: Facts) -> None:
    """Keep the deliberately failing evaluator result visible."""
    run = facts.demo_run("evaluation")
    rows = facts.evaluation_rows
    version = facts.packages.get("azure-ai-evaluation", "installed")
    slide = T.content_slide(
        prs,
        eyebrow="evaluation",
        title="Migration is only safe if the acceptance criteria survive",
        deck=f"azure-ai-evaluation {version} runs evaluators over a dataset. "
        "This recorded run uses Python checks, with a failing case "
        "kept visible.",
        number=33,
        accent=P.teal,
    )
    C.run_strip(slide, "evaluation")
    D.draw(
        slide,
        D.Schema(
            nodes=[
                D.Node(
                    "cases",
                    "cases.jsonl",
                    0,
                    0,
                    12,
                    3,
                    kind="state",
                    colour=P.blue,
                    mono=True,
                ),
                D.Node(
                    "sdk", "evaluate()", 17, 0, 12, 3, colour=P.teal, mono=True
                ),
                D.Node(
                    "checks",
                    "required_terms + source_sections",
                    34,
                    0,
                    17,
                    3,
                    colour=P.teal,
                ),
                D.Node("report", "rows + metrics", 56, 0, 10, 3, kind="note"),
            ],
            edges=[
                D.Edge("cases", "sdk", side="h"),
                D.Edge("sdk", "checks", side="h"),
                D.Edge("checks", "report", side="h"),
            ],
        ),
        x=T.MARGIN_L,
        y=T.Y_BODY,
    )
    C.table(
        slide,
        ("Case", "required_terms", "source_sections", "Verdict"),
        tuple(
            (
                str(row["case"]),
                "pass" if row["required_terms"]["passed"] else "fail",
                "pass" if row["source_sections"]["passed"] else "fail",
                "pass" if row["passed"] else "fail, on purpose",
            )
            for row in rows
        ),
        x=T.MARGIN_L,
        y=T.Y_BODY + 960120,
        w=T.COL_MAIN_W,
        widths=(2.6, 1.8, 1.8, 2),
        row_h=457200,
        size=15,
        accents=tuple(P.teal if row["passed"] else P.coral for row in rows),
    )
    metrics = run["sdk_metrics"]
    C.rail(
        slide,
        title="What evaluate() returned",
        lines=(
            *(f"{name}: {value:.2f}" for name, value in metrics.items()),
            f"failed_cases: {run['failed_cases']}",
            "Term and citation checks do not prove correctness.",
        ),
        y=T.Y_BODY + 777240,
        h=3200400,
        accent=P.teal,
    )
    _text(
        slide,
        f"mode={run['mode']}. These checks use no model and no "
        "project. Model-graded evaluation is a separate, paid and "
        "non-deterministic run against a deployed model.",
        y=int(6.7 * U),
        w=T.COL_MAIN_W,
        h=685800,
    )
    C.use_case(
        slide,
        who="QA engineer checks policy answers before release",
        steps=(
            "questions + expected sources",
            "run evaluators",
            "inspect failing cases",
        ),
        result="a regression report, with the failure kept visible.",
        x=T.MARGIN_L,
        y=int(7.85 * U),
        w=T.CONTENT_W,
        accent=P.teal,
    )
