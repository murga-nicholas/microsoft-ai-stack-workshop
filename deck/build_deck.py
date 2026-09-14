"""Build the workshop deck.

    uv run --no-sync python deck/build_deck.py

Every slide is a function. The order below is the running order, and
the numbers printed on the slides come from this list rather than from
anybody remembering to renumber.

The deck is built from the repository it documents: package versions
come from the installed environment, dated research has sourced
constants in ``research_facts.py``, and recorded evidence comes from
``deck/facts.json``, which ``collect_facts.py`` writes from a real run.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import TYPE_CHECKING

DECK_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(DECK_DIR))

import content_appendix  # noqa: E402
import content_close  # noqa: E402
import content_compare  # noqa: E402
import content_foundry  # noqa: E402
import content_legacy  # noqa: E402
import content_lineage  # noqa: E402
import content_map  # noqa: E402
import content_open  # noqa: E402
import content_runtime  # noqa: E402
import content_runtime_extra  # noqa: E402
import facts as facts_module  # noqa: E402
import theme as T  # noqa: E402
from appendix_code import attach_code_notes  # noqa: E402
from notes import attach_notes, load_notes  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

if TYPE_CHECKING:
    from collections.abc import Callable

    from facts import Facts

    SlideFn = Callable[[Presentation, Facts], None]

DEFAULT_OUTPUT = DECK_DIR.parent / "microsoft_ai_stack.pptx"
NOTES_FILE = DECK_DIR / "speaker_notes.md"
DEFAULT_DATE = "2026-09-14"
DEFAULT_AUTHOR = "Mykola Murha"
DEFAULT_ROLE = "Data and AI Engineer  ·  DataArt"

# The running order. Reordering this list changes the deck; renaming a
# function also requires updating its section key in the notes script.
SLIDES: tuple[SlideFn, ...] = (
    content_open.slide_title,
    content_open.slide_jobs,
    content_open.slide_agenda,
    content_lineage.slide_section,
    content_lineage.slide_framework_lineage,
    content_lineage.slide_service_lineage,
    content_lineage.slide_portal,
    content_lineage.slide_foundry_architecture,
    content_legacy.slide_section,
    content_legacy.slide_autogen,
    content_legacy.slide_semantic_kernel,
    content_legacy.slide_channel,
    content_legacy.slide_prompt_flow,
    content_runtime.slide_section,
    content_runtime.slide_agent,
    content_runtime_extra.slide_client_family,
    content_runtime.slide_workflow,
    content_runtime.slide_harness,
    content_runtime_extra.slide_mcp,
    content_runtime_extra.slide_a2a,
    content_runtime.slide_observability,
    content_compare.slide_section,
    content_compare.slide_comparison,
    content_compare.slide_recovery_evidence,
    content_foundry.slide_section,
    content_foundry.slide_models,
    content_foundry.slide_agent_service,
    content_foundry.slide_foundry_local,
    content_foundry.slide_local_comparison,
    content_foundry.slide_knowledge,
    content_foundry.slide_guardrails,
    content_foundry.slide_identity,
    content_foundry.slide_evaluation,
    content_map.slide_architectures,
    content_map.slide_decision,
    content_close.slide_close,
    content_appendix.slide_section,
    content_appendix.slide_transcripts,
    content_compare.slide_approval,
    content_runtime_extra.slide_mro,
    content_map.slide_stack_map,
    content_legacy.slide_cost_of_staying,
    content_map.slide_platform_decision,
    content_close.slide_operations,
    content_appendix.slide_repository,
    content_appendix.slide_readiness,
    content_appendix.slide_standard,
    content_appendix.slide_references,
)


# A slide is allowed to draw between the eyebrow and the footer. This
# is checked rather than trusted, because a table that overflows into
# the footer is exactly the kind of thing nobody notices until it is
# on a projector in front of a client.
SAFE_TOP = 457200
SAFE_BOTTOM = T.Y_FOOTER - 91440
SAFE_RIGHT = T.SLIDE_W - 457200

# Backgrounds legitimately cover the whole canvas.
FULL_BLEED = (T.SLIDE_W, T.SLIDE_H)


def check_layout(prs: Presentation) -> list[str]:
    """Return a complaint for every shape outside the safe area.

    Args:
        prs: The built presentation.

    Returns:
        Human-readable problems, empty when the deck is clean.
    """
    problems: list[str] = []
    for number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.width is None or shape.height is None:
                continue
            if (shape.width, shape.height) == FULL_BLEED:
                continue
            if str(shape.name).startswith("chrome-"):
                continue
            bottom = int(shape.top) + int(shape.height)
            right = int(shape.left) + int(shape.width)
            name = _describe(shape)
            if bottom > SAFE_BOTTOM:
                over = (bottom - SAFE_BOTTOM) / 914400
                problems.append(
                    f"slide {number}: {name} runs {over:.2f}in into the footer"
                )
            if right > SAFE_RIGHT:
                over = (right - SAFE_RIGHT) / 914400
                problems.append(
                    f"slide {number}: {name} runs {over:.2f}in off the "
                    f"right edge"
                )
            if int(shape.top) < SAFE_TOP:
                problems.append(f"slide {number}: {name} sits above the grid")
    return problems


def _describe(shape: object) -> str:
    """Name a shape by its text, falling back to its shape name."""
    text = ""
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        text = " ".join(text.split())[:40]
    return f"{text!r}" if text else str(getattr(shape, "name", "shape"))


# Arial and Consolas at the sizes this deck uses average out at roughly
# this fraction of the point size per character. It is an estimate, not
# a font metric, so the tolerance below is deliberately generous: the
# goal is to catch a paragraph that overruns its box by a line or more,
# not to police kerning.
CHAR_WIDTH_RATIO = 0.52
FIT_TOLERANCE = 1.12

# PowerPoint's "single" line is taller than the point size: about 1.15
# of it for Arial and 1.17 for Consolas. Without this factor a code
# panel of a dozen lines passes the estimate and still spills out of
# its panel on the projector.
LINE_HEIGHT = 1.16

# A single-line label that is a hair taller than its box renders fine,
# because text boxes do not clip. Only an overrun of a tenth of an inch
# or more - a partial line - is worth failing the build for.
MIN_OVERRUN = 91440


def check_text_fit(prs: Presentation) -> list[str]:
    """Return a complaint for every text box whose text will clip.

    Args:
        prs: The built presentation.

    Returns:
        Human-readable problems, empty when every box holds its text.
    """
    problems: list[str] = []
    for number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if str(shape.name).startswith("chrome-"):
                continue
            if not shape.text_frame.text.strip():
                continue
            if shape.name == "title":
                problems.extend(_title_problems(number, shape))
                continue
            needed = _estimated_height(shape)
            available = int(shape.height) * FIT_TOLERANCE
            if needed - available > MIN_OVERRUN:
                over = (needed - available) / 914400
                problems.append(
                    f"slide {number}: {_describe(shape)} needs "
                    f"{over:.2f}in more height than its box"
                )
    return problems


# Arial advance widths in thousandths of an em. A wrapped headline
# collides with the deck line below it, so headlines are measured with
# real glyph widths rather than the average used for body text.
ARIAL_WIDTHS = (
    dict(
        zip(
            "abcdefghijklmnopqrstuvwxyz",
            (
                556,
                556,
                500,
                556,
                556,
                278,
                556,
                556,
                222,
                222,
                500,
                222,
                833,
                556,
                556,
                556,
                556,
                333,
                500,
                278,
                556,
                500,
                722,
                500,
                500,
                500,
            ),
            strict=True,
        )
    )
    | dict(
        zip(
            "ABCDEFGHIJKLMNOPQRSTUVWXYZ",
            (
                667,
                667,
                722,
                722,
                667,
                611,
                778,
                722,
                278,
                500,
                667,
                556,
                833,
                722,
                778,
                667,
                778,
                722,
                667,
                611,
                722,
                667,
                944,
                667,
                667,
                611,
            ),
            strict=True,
        )
    )
    | {" ": 278, ",": 278, ".": 278, ";": 278, ":": 278, "-": 333}
)
ARIAL_DEFAULT_WIDTH = 556


def _title_problems(number: int, shape: object) -> list[str]:
    """Complain when a slide headline will not fit on one line."""
    text = shape.text_frame.text.strip()
    size = shape.text_frame.paragraphs[0].runs[0].font.size.pt
    width = sum(ARIAL_WIDTHS.get(char, ARIAL_DEFAULT_WIDTH) for char in text)
    needed = width / 1000 * size * 12700
    if needed <= int(shape.width):
        return []
    over = (needed - int(shape.width)) / 914400
    return [
        f"slide {number}: headline {text[:40]!r} is {over:.2f}in too "
        f"wide for one line"
    ]


def _estimated_height(shape: object) -> float:
    """Estimate the rendered height of a text box, in EMU."""
    width = int(shape.width)
    total = 0.0
    for para in shape.text_frame.paragraphs:
        runs = [run for run in para.runs if run.text]
        if not runs:
            continue
        size = max(
            (run.font.size.pt for run in runs if run.font.size),
            default=14.0,
        )
        characters = sum(len(run.text) for run in runs)
        per_line = max(
            int(width / (size * CHAR_WIDTH_RATIO * 12700)),
            1,
        )
        lines = max(1, -(-characters // per_line))
        spacing = para.line_spacing or 1.2
        spacing = spacing if isinstance(spacing, float) else 1.2
        total += lines * size * spacing * LINE_HEIGHT * 12700
        if para.space_after is not None:
            total += para.space_after.pt * 12700
    return total


def build(output: Path, facts: Facts, *, draft: bool = False) -> Path:
    """Render every slide and save the presentation.

    Args:
        output: Where to write the ``.pptx``.
        facts: Numbers the slides are allowed to quote.
        draft: Save even when validation checks fail, so the problems
            can be looked at in a render. Never ship a draft build.

    Returns:
        The path written.
    """
    notes_problems: list[str] = []
    scripts: dict[str, str] = {}
    try:
        scripts = load_notes(NOTES_FILE)
    except FileNotFoundError:
        notes_problems.append(f"notes file missing: {NOTES_FILE}")
    except ValueError as error:
        notes_problems.append(f"invalid notes file {NOTES_FILE}: {error}")
    else:
        keys = [f"{render.__module__}.{render.__name__}" for render in SLIDES]
        notes_problems.extend(
            f"missing notes section: {key}"
            for key in keys
            if key not in scripts
        )
        notes_problems.extend(
            f"stale notes section (not in SLIDES): {key}"
            for key in sorted(scripts.keys() - set(keys))
        )

    prs = Presentation()
    prs.slide_width = Emu(T.SLIDE_W)
    prs.slide_height = Emu(T.SLIDE_H)

    T.register_order([render.__name__ for render in SLIDES])
    for index, render in enumerate(SLIDES, start=1):
        T.begin_slide(index)
        try:
            count_before = len(prs.slides)
            render(prs, facts)
            assert len(prs.slides) == count_before + 1, (  # noqa: S101
                "each slide function must add exactly one slide"
            )
            key = f"{render.__module__}.{render.__name__}"
            if key in scripts:
                attach_notes(prs.slides[-1], scripts[key])
        except Exception as error:
            message = f"slide {index} ({render.__name__}) failed: {error}"
            raise RuntimeError(message) from error

    attach_code_notes(prs)
    problems = check_layout(prs) + check_text_fit(prs)
    messages: list[str] = []
    if problems:
        joined = "\n  ".join(problems)
        messages.append(f"layout is not clean:\n  {joined}")
    if notes_problems:
        joined = "\n  ".join(notes_problems)
        messages.append(f"speaker notes are not valid:\n  {joined}")
    joined = "\n".join(messages)
    if messages and not draft:
        message = f"deck {joined}"
        raise RuntimeError(message)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    if messages:
        print(f"DRAFT - {joined}", file=sys.stderr)
    return output


def main(argv: list[str] | None = None) -> int:
    """Build the deck from the command line."""
    parser = argparse.ArgumentParser(
        prog="build_deck",
        description="Build the Microsoft AI stack workshop deck.",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--date", default=DEFAULT_DATE)
    parser.add_argument("--author", default=DEFAULT_AUTHOR)
    parser.add_argument("--role", default=DEFAULT_ROLE)
    parser.add_argument(
        "--draft",
        action="store_true",
        help="save despite layout or notes problems, for inspection only",
    )
    args = parser.parse_args(argv)

    facts = facts_module.load(
        deck_date=args.date,
        author=args.author,
        author_role=args.role,
    )
    written = build(args.output, facts, draft=args.draft)
    print(f"{len(SLIDES)} slides -> {written}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
